"""Rebuild Tim's price sheet from the pristine blank template.

Source: reference/JQC-EXCEL-PRICE-SHEET-1.xlsx  (untouched upload, never written to)
Output: Next Level Excel.xlsx                    (Tim's working sheet)

What it does:
  - writes the '#Section' heading rows the proposal builder will key off
  - adds the customer detail fields next to Job Name
  - widens the Carpentry and Plumbing line bands
  - repairs every formula whose row moved (openpyxl does not rewrite them)

Refuses to overwrite an existing output unless --force is given, so numbers
Tim has already typed in can never be clobbered by accident.
"""

import os
import sys
from copy import copy

import openpyxl

from lock_price_sheet import apply_lock

SRC = os.path.join("reference", "JQC-EXCEL-PRICE-SHEET-1.xlsx")
OUT = "Next Level Excel.xlsx"

CARPENTRY_EXTRA = 10  # 5 blank lines -> 15
PLUMBING_EXTRA = 4    # 4 blank lines -> 8

LINE_HEADINGS = {5: "#Demolition", 10: "#Carpentry", 26: "#Plumbing", 35: "#Electrical"}
SUBS_HEADINGS = {2: "#Plumbing", 6: "#Electrical", 10: "#Drywall", 14: "#Masonry", 18: "#Painting"}
INSERTED_ROWS = list(range(16, 26)) + list(range(31, 35))
# every row from 5 to 43 is a line row - headings included: they carry the same
# Our Rate / Our Total Cost / Retail Rate / Customer Total Price formulas
LINE_ROWS = list(range(5, 44))
CREW_ROWS = list(range(44, 51))
LABOR_TOTAL_ROW = 51


def main():
    if os.path.exists(OUT) and "--force" not in sys.argv:
        sys.exit("%s already exists - pass --force to overwrite it." % OUT)

    wb = openpyxl.load_workbook(SRC)
    labor = wb["Labor Estimate"]
    summary = wb["Summary Dashboard"]

    # --- widen the two bands -------------------------------------------------
    labor.insert_rows(16, CARPENTRY_EXTRA)  # room for Carpentry
    labor.insert_rows(31, PLUMBING_EXTRA)   # room for Plumbing

    # inserted rows arrive with no formatting - borrow it from a normal line row
    for r in INSERTED_ROWS:
        for c in range(1, 7):
            labor.cell(row=r, column=c)._style = copy(labor.cell(row=6, column=c)._style)
        labor.row_dimensions[r].height = labor.row_dimensions[6].height

    # --- section headings ----------------------------------------------------
    for row, label in LINE_HEADINGS.items():
        labor.cell(row=row, column=1).value = label

    # --- subcontractor section headings -------------------------------------
    for row, label in SUBS_HEADINGS.items():
        wb["Subs"].cell(row=row, column=1).value = label

    # --- labor line formulas -------------------------------------------------
    for r in LINE_ROWS:
        labor.cell(row=r, column=3).value = "=$F$2"
        labor.cell(row=r, column=4).value = "=IF(B{r}>0, B{r}*C{r}, 0)".format(r=r)
        labor.cell(row=r, column=5).value = "=$F$3"
        labor.cell(row=r, column=6).value = "=IF(B{r}>0, B{r}*E{r}, 0)".format(r=r)

    # --- 3-man crew block (its rates point at the block's own top row) -------
    labor.cell(row=CREW_ROWS[0], column=1).value = (
        "3-Man Crew Premium Upgrade (Lead Carpenter + 2 Helpers)"
    )
    labor.cell(row=CREW_ROWS[0], column=3).value = 73
    labor.cell(row=CREW_ROWS[0], column=5).value = 120
    for r in CREW_ROWS:
        labor.cell(row=r, column=4).value = "=IF(B{r}>0, B{r}*C{r}, 0)".format(r=r)
        labor.cell(row=r, column=6).value = "=IF(B{r}>0, B{r}*E{r}, 0)".format(r=r)
    for r in CREW_ROWS[1:]:
        labor.cell(row=r, column=3).value = "=$C$%d" % CREW_ROWS[0]
        labor.cell(row=r, column=5).value = "=$E$%d" % CREW_ROWS[0]

    # --- labor total ---------------------------------------------------------
    labor.cell(row=LABOR_TOTAL_ROW, column=1).value = "Total Labor"
    labor.cell(row=LABOR_TOTAL_ROW, column=2).value = "=SUM(B5:B%d)" % (LABOR_TOTAL_ROW - 1)
    labor.cell(row=LABOR_TOTAL_ROW, column=4).value = "=SUM(D5:D%d)" % (LABOR_TOTAL_ROW - 1)
    labor.cell(row=LABOR_TOTAL_ROW, column=6).value = "=SUM(F5:F%d)" % (LABOR_TOTAL_ROW - 1)

    # --- customer fields on the summary sheet --------------------------------
    summary.insert_rows(6, 3)
    for r in range(6, 9):
        for c in range(1, 6):
            summary.cell(row=r, column=c)._style = copy(summary.cell(row=5, column=c)._style)
    summary["A6"].value = "Customer Name:"
    summary["A7"].value = "Customer Address:"
    summary["A8"].value = "Customer Phone:"

    # --- summary formulas (every row moved down by 3) ------------------------
    summary["B11"] = "='Materials Estimate'!D31"
    summary["C11"] = "='Materials Estimate'!F31"
    summary["D11"] = "=C11-B11"
    summary["E11"] = "=IF(C11>0, D11/C11, 0)"
    summary["B12"] = "='Labor Estimate'!D%d" % LABOR_TOTAL_ROW
    summary["C12"] = "='Labor Estimate'!F%d" % LABOR_TOTAL_ROW
    summary["D12"] = "=C12-B12"
    summary["E12"] = "=IF(C12>0, D12/C12, 0)"
    summary["B13"] = "=Subs!D25"
    summary["C13"] = "=Subs!F25"
    summary["D13"] = "=C13-B13"
    summary["E13"] = "=IF(C13>0, D13/C13, 0)"
    summary["B14"] = "=B11+B12+B13"
    summary["C14"] = "=C11+C12+C13"
    summary["D14"] = "=C14-B14"
    summary["E14"] = "=IF(C14>0, D14/C14, 0)"
    summary["C15"] = "=C14*0.06"
    summary["D15"] = "=C15-B15"
    summary["B16"] = "=B14+B15"
    summary["C16"] = "=C14+C15"
    summary["D16"] = "=C16-B16"
    summary["E16"] = "=IF(C16>0, D16/C16, 0)"
    summary["C19"] = "=C16"
    summary["C20"] = "=C19*0.5"
    summary["C21"] = "=C19*0.25"
    summary["C22"] = "=C19*0.25"
    summary["C23"] = "=C20+C21+C22"

    # portrait + label lock, so a rebuild can never hand Tim back a landscape,
    # wide-open sheet
    apply_lock(wb)

    wb.save(OUT)
    print("wrote", OUT)


main()
