"""Throwaway check: after the '#' cleanup, are the rows, formulas and saved totals still there?"""
import os

import openpyxl

BASE = r"C:\Users\user\Documents\July's Quality Construction"
TARGETS = [
    'JQC templets\\Next_Level_Excel_MASTER.xlsx',
    'Kerry Thompson Estimate\\Kerry_Thompson_Kitchen_20_showroom_lines.xlsx',
]
TABS = ('Labor Estimate', 'Subs', 'Summary Dashboard')


def show(path):
    print('===', os.path.basename(path))
    formulas = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    print('  tabs:', formulas.sheetnames)
    for tab in TABS:
        if tab not in formulas.sheetnames:
            print(f'  (no {tab} tab)')
            continue
        ws = formulas[tab]
        wsv = values[tab]
        print(f'  {tab}: {ws.max_row} rows x {ws.max_column} cols')
        for row in range(1, ws.max_row + 1):
            label = ws.cell(row=row, column=1).value
            money = {
                ws.cell(row=row, column=col).coordinate: (
                    f'{ws.cell(row=row, column=col).value} -> {wsv.cell(row=row, column=col).value}'
                )
                for col in (3, 4, 6, 7)
                if ws.cell(row=row, column=col).value is not None
            }
            if label is not None or money:
                print(f'    row {row}: A={label!r} {money if money else ""}')
    print()


for rel in TARGETS:
    show(os.path.join(BASE, rel))
