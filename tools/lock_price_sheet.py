"""Locks Tim's price sheet down and pins it to portrait.

Every cell on every tab is set to "locked" except the blank boxes Tim types in:
the description / quantity / unit cost / rate boxes on Materials, Labor and Subs,
and the boxes beside Job Name, Date, Customer Name, Customer Address and Customer
Phone on the Summary Dashboard.

Locked means: labels, section headings, totals, rates, the tax rate, markup % and
every formula. A stray keystroke cannot wipe them.

Sheet protection goes on with NO password. It is there to catch a slip, not to
keep anybody out, so it can be switched off again from the same LibreOffice menu
it was switched on from (Tools > Protect Sheet).

Usage: python tools/lock_price_sheet.py ["path to workbook"]
"""

import sys

from openpyxl import load_workbook
from openpyxl.styles import Protection
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.protection import SheetProtection

DEFAULT = "Next Level Excel.xlsx"

MATERIAL_ROWS = list(range(5, 30)) + [30]
SHOWROOM_ROWS = list(range(36, 46))

LABOR_LINE_ROWS = list(range(5, 44))
LABOR_HEADING_ROWS = {5, 10, 26, 35}
LABOR_CREW_ROWS = list(range(44, 51))

SUBS_ROWS = list(range(2, 24))
SUBS_HEADING_ROWS = {2, 6, 10, 14, 18}

SUMMARY_INPUT_BOXES = {
    (5, 2),  # Job Name
    (5, 5),  # Date: label sits in D5, so its box is E5
    (6, 2),  # Customer Name
    (7, 2),  # Customer Address
    (8, 2),  # Customer Phone
}


def typing_boxes():
    """(sheet, row, column) for every cell Tim is meant to be able to type in."""
    boxes = set()

    for row, col in SUMMARY_INPUT_BOXES:
        boxes.add(("Summary Dashboard", row, col))

    # Materials: item, qty, unit cost (+ the dump fee row); showroom also takes a
    # retail price per unit in column E.
    for row in MATERIAL_ROWS:
        for col in (1, 2, 3):
            boxes.add(("Materials Estimate", row, col))
    for row in SHOWROOM_ROWS:
        for col in (1, 2, 3, 5):
            boxes.add(("Materials Estimate", row, col))

    # Labor: description and hours. The #Section headings are labels, not boxes.
    for row in LABOR_LINE_ROWS:
        cols = (2,) if row in LABOR_HEADING_ROWS else (1, 2)
        for col in cols:
            boxes.add(("Labor Estimate", row, col))
    for row in LABOR_CREW_ROWS:
        boxes.add(("Labor Estimate", row, 2))
    for row in LABOR_CREW_ROWS[1:]:
        boxes.add(("Labor Estimate", row, 1))

    # Subs: trade, qty, their cost per unit.
    for row in SUBS_ROWS:
        cols = (2, 3) if row in SUBS_HEADING_ROWS else (1, 2, 3)
        for col in cols:
            boxes.add(("Subs", row, col))

    return boxes


def pin_portrait(ws):
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def apply_lock(wb):
    boxes = typing_boxes()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                open_for_typing = (ws.title, cell.row, cell.column) in boxes
                cell.protection = Protection(locked=not open_for_typing)
        # A fresh protection object, deliberately: the blank template arrived
        # already protected with a password we do not have. Building our own
        # leaves no password on the file, so LibreOffice's Protect Sheet menu
        # toggles the lock off without ever asking for one.
        ws.protection = SheetProtection(sheet=True)
        pin_portrait(ws)


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    wb = load_workbook(path)
    apply_lock(wb)
    wb.save(path)
    print("locked and pinned portrait:", path)


if __name__ == "__main__":
    main()
