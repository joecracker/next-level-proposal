"""Read-only check on the lock: portrait everywhere, protection on, and only the
blank typing boxes open.

Usage: python tools/verify_lock.py ["path to workbook"]
"""

import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from lock_price_sheet import typing_boxes

DEFAULT = "Next Level Excel.xlsx"

EXPECTED_BOXES = {
    "Summary Dashboard": "B5,E5,B6,B7,B8",
    "Materials Estimate": "A/B/C on rows 5-30 and 36-45, plus E on 36-45",
    "Labor Estimate": "A/B on the line rows, B on the crew rows, A on 45-50",
    "Subs": "A/B/C on rows 2-23, minus the #Section heading cells",
}


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    wb = load_workbook(path)

    for ws in wb.worksheets:
        open_cells = [
            "%s%d" % (get_column_letter(c.column), c.row)
            for row in ws.iter_rows()
            for c in row
            if c.protection.locked is False
        ]
        print("---", ws.title)
        print("  orientation:", ws.page_setup.orientation)
        print("  fit to page:", ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage,
              "/ width", ws.page_setup.fitToWidth, "height", ws.page_setup.fitToHeight)
        print("  protection on:", ws.protection.sheet, "/ password:", ws.protection.password)
        print("  open boxes:", len(open_cells))
        print("  sample:", ", ".join(open_cells[:8]))

    expected = typing_boxes()
    found = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.protection.locked is False:
                    found.add((ws.title, c.row, c.column))

    missing = sorted(expected - found)
    extra = sorted(found - expected)
    print("\nboxes expected but locked:", missing or "none")
    print("cells open that shouldn't be:", extra or "none")

    # the labels and totals must still be there and intact
    checks = [
        ("Summary Dashboard", "A5", "Job Name:"),
        ("Summary Dashboard", "D5", "Date:"),
        ("Summary Dashboard", "A6", "Customer Name:"),
        ("Summary Dashboard", "A7", "Customer Address:"),
        ("Summary Dashboard", "A8", "Customer Phone:"),
        ("Summary Dashboard", "A19", "Total Contract Amount"),
        ("Labor Estimate", "A5", "#Demolition"),
        ("Labor Estimate", "A10", "#Carpentry"),
        ("Labor Estimate", "A51", "Total Labor"),
        ("Subs", "A2", "#Plumbing"),
        ("Subs", "A25", "Total Sub Cost"),
        ("Materials Estimate", "A31", "Total Materials & Fees"),
    ]
    print()
    for sheet, ref, want in checks:
        got = wb[sheet][ref].value
        flag = "ok" if got == want else "CHANGED"
        print("%-6s %s!%s = %r" % (flag, sheet, ref, got))


if __name__ == "__main__":
    main()
